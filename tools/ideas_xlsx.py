#!/usr/bin/env python3
"""靈感整理 Excel — rebuilds ideas/靈感整理.xlsx from IDEAS.md (his call, 2026-10-29: "it's a list").

IDEAS.md is still the only place entries are typed; this reads it with the same parser as
tools/ideas_doc.py and writes two sheets:

  目錄  one line per 類型 with its count; click a type to jump to its rows.
  靈感  every entry, grouped by 類型 in the usual order, one row each: 類型, 名稱, 重點, 狀態,
        喜歡它什麼, 日期, 連結. The header is frozen and filterable. An entry with two types
        (類型: 網站靈感 / UI 介面靈感) gets a row under each, so filtering by either finds it.

Run from anywhere:  python3 tools/ideas_xlsx.py   (python3 tools/ideas_doc.py with no arguments
does the same). Safe to rerun: the file is overwritten. Never edit the Excel file by hand.
The build list (TODO.md) and the weekly reports still go to Word through tools/ideas_doc.py.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from ideas_doc import (FRAME, HEAD, IDEAS, INK, KODAK_RED, MUTED, SOURCE, by_type, kinds,
                       parse, point, url)

OUT = IDEAS / '靈感整理.xlsx'
FONT = 'Arial'
# column title, width, field
COLUMNS = [('類型', 14, None), ('名稱', 30, None), ('重點', 34, None), ('狀態', 9, '狀態'),
           ('喜歡它什麼', 44, '喜歡它什麼'), ('日期', 12, '日期'), ('連結', 7, None)]

line = Side(style='thin', color=FRAME)
FRAMED = Border(left=line, right=line, top=line, bottom=line)
HEAD_FILL = PatternFill('solid', fgColor=HEAD)
TOP = Alignment(vertical='top', wrap_text=True)


def font(**kw):
    return Font(name=FONT, size=kw.pop('size', 10), color=kw.pop('color', INK), **kw)


def header(ws, row, titles):
    for col, title in enumerate(titles, 1):
        cell = ws.cell(row, col, title)
        cell.font, cell.fill, cell.border = font(bold=True), HEAD_FILL, FRAMED
        cell.alignment = Alignment(vertical='center')


def main(source=SOURCE, out=OUT):
    entries = parse(source.read_text(encoding='utf-8'))
    groups = by_type(entries)
    wb = Workbook()

    toc = wb.active
    toc.title = '目錄'
    sheet = wb.create_sheet('靈感')

    # 靈感: one row per entry per type
    header(sheet, 1, [c[0] for c in COLUMNS])
    first_row, row = {}, 2
    for kind, found in groups.items():
        first_row[kind] = row
        for entry in found:
            fields = entry['fields']
            values = [kind, entry['title'], point(fields)]
            values += [fields.get(c[2], '') for c in COLUMNS[3:6]]
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row, col, value)
                cell.font, cell.border, cell.alignment = font(), FRAMED, TOP
            sheet.cell(row, 1).font = font(color=MUTED)
            sheet.cell(row, 2).font = font(bold=True)
            link = sheet.cell(row, len(COLUMNS))
            link.border, link.alignment = FRAMED, TOP
            if url(fields):
                link.value, link.hyperlink = '開啟', url(fields)
                link.font = font(color=KODAK_RED, underline='single')
            row += 1
    for col, (_, width, _) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{row - 1}'

    # 目錄: title, one line about the file, then each type and its count
    toc['A1'] = '靈感整理'
    toc['A1'].font = font(size=16, bold=True)
    both = sum(len(kinds(e)) > 1 for e in entries)
    toc['A2'] = (f'最後更新 {date.today().isoformat()}  ·  {len(entries)} 項，{len(groups)} 個類型'
                 + (f'（{both} 項同時屬兩類，兩邊都列）' if both else '')
                 + f'  ·  來源 {source.name}')
    toc['A2'].font = font(color=MUTED)
    toc['A3'] = '點類型跳到那一類；在「靈感」表裡也可以用標題列的篩選鈕挑類型或狀態。'
    toc['A3'].font = font(color=MUTED)
    header(toc, 5, ['類型', '數量'])
    for n, (kind, found) in enumerate(groups.items(), 6):
        name, count = toc.cell(n, 1, kind), toc.cell(n, 2, len(found))
        name.hyperlink = Hyperlink(ref=name.coordinate, location=f'靈感!A{first_row[kind]}')
        name.font = font(color=KODAK_RED, underline='single')
        count.font = font()
        name.border = count.border = FRAMED
    toc.column_dimensions['A'].width = 24
    toc.column_dimensions['B'].width = 8

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f'{out} — {len(entries)} 項, {len(groups)} 個類型')


if __name__ == '__main__':
    main()
